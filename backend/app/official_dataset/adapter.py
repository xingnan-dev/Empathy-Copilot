"""Read-only official Excel adapter: read, normalize, and validate only."""

import re
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from backend.app.domain.scenarios import scenario_from_business_scene

from .models import (
    ConversationSession,
    MessageRecord,
    OfficialConsumerRecord,
    OfficialDataset,
    OfficialOrderRecord,
    OfficialServiceRecord,
)


class OfficialDatasetValidationError(ValueError):
    pass


CHAT_HEADERS = {
    "会话ID",
    "消息序号",
    "message_id",
    "发送时间",
    "角色",
    "买家昵称",
    "发送方",
    "店铺",
    "scene_major",
    "scene_minor",
    "message_text",
    "内容类型",
    "image_path",
    "关联订单号",
    "关联工单号",
}
ORDER_HEADERS = {
    "订单号",
    "会话ID",
    "买家昵称",
    "商品货号",
    "商品名称",
    "数量",
    "单价(元)",
    "实付金额(元)",
    "订单状态",
    "下单时间",
    "付款时间",
    "发货时间",
    "快递公司",
    "物流单号",
}
SERVICE_SHEETS = {
    "补发换货工单": ("工单状态", "售后原因"),
    "线下打款工单": ("工单状态", "退款问题类型"),
    "物流工单": ("工单状态", "问题类型"),
    "不良反应工单": ("任务状态", "类型"),
    "售后退货工单": ("任务状态", "退货原因"),
}


class OfficialDatasetAdapter:
    def read(self, path: str | Path) -> OfficialDataset:
        source = Path(path)
        if not source.is_file():
            raise OfficialDatasetValidationError(f"official dataset not found: {source.name}")
        workbook = load_workbook(source, read_only=True, data_only=True)
        chat = self._rows(workbook["聊天记录"], CHAT_HEADERS)
        orders = self._rows(workbook["订单"], ORDER_HEADERS)
        service_rows = {
            name: self._rows(
                workbook[name],
                {"工单号", "会话ID", "关联订单号", "买家昵称", status, issue, "创建时间"},
            )
            for name, (status, issue) in SERVICE_SHEETS.items()
        }
        return self._normalize(chat, orders, service_rows)

    @staticmethod
    def _rows(sheet: Any, required: set[str]) -> list[dict[str, Any]]:
        iterator = sheet.iter_rows(values_only=True)
        try:
            headers = [str(value) if value is not None else "" for value in next(iterator)]
        except StopIteration as error:
            raise OfficialDatasetValidationError(f"sheet {sheet.title} is empty") from error
        missing = required - set(headers)
        if missing:
            raise OfficialDatasetValidationError(
                f"sheet {sheet.title} missing required columns: {sorted(missing)}"
            )
        return [
            dict(zip(headers, values, strict=True))
            for values in iterator
            if any(value is not None for value in values)
        ]

    def _normalize(
        self,
        chat: list[dict[str, Any]],
        orders: list[dict[str, Any]],
        service_sheets: dict[str, list[dict[str, Any]]],
    ) -> OfficialDataset:
        messages = [self._message(row) for row in chat]
        if len({item.message_id for item in messages}) != len(messages):
            raise OfficialDatasetValidationError("message_id must be unique")
        normalized_orders = [self._order(row) for row in orders]
        if len({item.order_id for item in normalized_orders}) != len(normalized_orders):
            raise OfficialDatasetValidationError("order_id must be unique")
        records = [
            self._service(sheet, row) for sheet, rows in service_sheets.items() for row in rows
        ]
        if len({item.service_record_id for item in records}) != len(records):
            raise OfficialDatasetValidationError("service_record_id must be unique")
        sessions = self._sessions(chat)
        consumers = [
            OfficialConsumerRecord(consumer_id=value)
            for value in sorted({item.consumer_id for item in sessions})
        ]
        self._validate_links(sessions, messages, normalized_orders, records)
        return OfficialDataset(
            consumers=consumers,
            sessions=sessions,
            messages=messages,
            orders=normalized_orders,
            service_records=records,
        )

    @staticmethod
    def _message(row: dict[str, Any]) -> MessageRecord:
        return MessageRecord(
            message_id=str(row["message_id"]),
            session_id=str(row["会话ID"]),
            sequence=int(row["消息序号"]),
            sent_at=OfficialDatasetAdapter._datetime(row["发送时间"]),
            consumer_id=str(row["买家昵称"]),
            role=str(row["角色"]),
            sender=str(row["发送方"]),
            text=str(row["message_text"]),
            content_type=str(row["内容类型"]),
            image_path=OfficialDatasetAdapter._optional(row.get("image_path")),
            order_id=OfficialDatasetAdapter._optional(row.get("关联订单号")),
            ticket_id=OfficialDatasetAdapter._optional(row.get("关联工单号")),
        )

    @staticmethod
    def _order(row: dict[str, Any]) -> OfficialOrderRecord:
        return OfficialOrderRecord(
            order_id=str(row["订单号"]),
            session_id=str(row["会话ID"]),
            consumer_id=str(row["买家昵称"]),
            product_id=str(row["商品货号"]),
            product_name=str(row["商品名称"]),
            quantity=int(row["数量"]),
            unit_price=float(row["单价(元)"]),
            paid_amount=float(row["实付金额(元)"]),
            order_status=str(row["订单状态"]),
            ordered_at=OfficialDatasetAdapter._datetime(row["下单时间"]),
            paid_at=OfficialDatasetAdapter._datetime(row["付款时间"]),
            shipped_at=OfficialDatasetAdapter._optional_datetime(row.get("发货时间")),
            carrier=OfficialDatasetAdapter._optional(row.get("快递公司")),
            tracking_number=OfficialDatasetAdapter._optional(row.get("物流单号")),
        )

    @staticmethod
    def _service(sheet: str, row: dict[str, Any]) -> OfficialServiceRecord:
        status_column, issue_column = SERVICE_SHEETS[sheet]
        excluded = {"支付宝账号", "支付宝实名", "处理人"}
        details = {
            str(key): OfficialDatasetAdapter._scalar(value)
            for key, value in row.items()
            if key not in excluded
            and key
            not in {
                "工单号",
                "会话ID",
                "关联订单号",
                "买家昵称",
                status_column,
                issue_column,
                "创建时间",
                "完成时间",
            }
        }
        return OfficialServiceRecord(
            service_record_id=str(row["工单号"]),
            source_kind=sheet,
            session_id=str(row["会话ID"]),
            order_id=str(row["关联订单号"]),
            consumer_id=str(row["买家昵称"]),
            issue_type=str(row[issue_column]),
            status=str(row[status_column]),
            created_at=OfficialDatasetAdapter._datetime(row["创建时间"]),
            completed_at=OfficialDatasetAdapter._optional_datetime(row.get("完成时间")),
            details=details,
        )

    @staticmethod
    def _sessions(chat: list[dict[str, Any]]) -> list[ConversationSession]:
        grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
        for row in chat:
            grouped[str(row["会话ID"])].append(row)
        sessions = []
        for session_id, rows in grouped.items():
            first = min(rows, key=lambda item: OfficialDatasetAdapter._datetime(item["发送时间"]))
            sessions.append(
                ConversationSession(
                    session_id=session_id,
                    consumer_id=str(first["买家昵称"]),
                    shop=str(first["店铺"]),
                    scene_major=str(first["scene_major"]),
                    scene_minor=str(first["scene_minor"]),
                    scenario=scenario_from_business_scene(str(first["scene_major"])),
                    started_at=OfficialDatasetAdapter._datetime(first["发送时间"]),
                    order_ids=sorted(
                        {str(row["关联订单号"]) for row in rows if row.get("关联订单号")}
                    ),
                    ticket_ids=sorted(
                        {str(row["关联工单号"]) for row in rows if row.get("关联工单号")}
                    ),
                )
            )
        return sorted(sessions, key=lambda item: item.session_id)

    @staticmethod
    def _validate_links(
        sessions: list[ConversationSession],
        messages: list[MessageRecord],
        orders: list[OfficialOrderRecord],
        records: list[OfficialServiceRecord],
    ) -> None:
        session_ids = {item.session_id for item in sessions}
        linked_session_ids = [item.session_id for item in messages]
        linked_session_ids.extend(item.session_id for item in orders)
        linked_session_ids.extend(item.session_id for item in records)
        if any(session_id not in session_ids for session_id in linked_session_ids):
            raise OfficialDatasetValidationError("record references an unknown session")
        order_ids = {item.order_id for item in orders}
        if any(item.order_id not in order_ids for item in records):
            raise OfficialDatasetValidationError("service record references an unknown order")

    @staticmethod
    def _datetime(value: Any) -> datetime:
        if isinstance(value, datetime):
            return value
        text = str(value).strip()
        match = re.fullmatch(r"(\d{4}-\d{2}-\d{2}[ T]\d{2}:\d{2}:\d{2}).*", text)
        try:
            return datetime.fromisoformat(match.group(1) if match else text)
        except ValueError as error:
            raise OfficialDatasetValidationError(f"invalid datetime: {value!r}") from error

    @staticmethod
    def _optional_datetime(value: Any) -> datetime | None:
        return None if value in {None, ""} else OfficialDatasetAdapter._datetime(value)

    @staticmethod
    def _optional(value: Any) -> str | None:
        return None if value in {None, ""} else str(value)

    @staticmethod
    def _scalar(value: Any) -> str | float | int | bool | None:
        if value is None or isinstance(value, str | float | int | bool):
            return value
        if isinstance(value, datetime):
            return value.isoformat()
        return str(value)
